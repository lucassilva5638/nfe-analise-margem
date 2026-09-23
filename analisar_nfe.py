"""
Lê uma pasta de XMLs de NF-e e gera um relatório gerencial em Excel:
DRE mensal, margem por produto, curva ABC e checagem de qualidade dos dados.

Uso:  python analisar_nfe.py --pasta xmls --cnpj 11222333000181 --saida relatorio_margem.xlsx
"""
import argparse, glob, os
import xml.etree.ElementTree as ET
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

NS = {"n": "http://www.portalfiscal.inf.br/nfe"}
MESES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]

# ---------------------------------------------------------------- leitura
def txt(el, path, default=""):
    x = el.find(path, NS)
    return x.text if x is not None and x.text is not None else default

def num(el, path):
    return float(txt(el, path, "0") or 0)

def dv_ok(chave):
    if len(chave) != 44 or not chave.isdigit(): return False
    pesos, soma = [2, 3, 4, 5, 6, 7, 8, 9], 0
    for i, d in enumerate(reversed(chave[:43])): soma += int(d) * pesos[i % 8]
    r = soma % 11
    return chave[43] == ("0" if r < 2 else str(11 - r))

def ler_xmls(pasta, cnpj):
    itens, ocorrencias, vistos = [], [], set()
    arquivos = sorted(glob.glob(os.path.join(pasta, "*.xml")))
    for arq in arquivos:
        nome = os.path.basename(arq)
        try:
            raiz = ET.parse(arq).getroot()
        except ET.ParseError:
            ocorrencias.append((nome, "", "XML ilegível", "Arquivo ignorado")); continue
        inf = raiz.find(".//n:infNFe", NS)
        if inf is None:
            ocorrencias.append((nome, "", "Não é NF-e", "Arquivo ignorado")); continue
        chave = inf.get("Id", "")[3:]
        cstat = txt(raiz, ".//n:protNFe/n:infProt/n:cStat")
        if chave in vistos:
            ocorrencias.append((nome, chave, "Nota duplicada", "Mantida só a primeira")); continue
        vistos.add(chave)
        if not dv_ok(chave):
            ocorrencias.append((nome, chave, "Dígito verificador da chave inválido", "Nota ignorada — conferir com o emitente")); continue
        if cstat != "100":
            ocorrencias.append((nome, chave, f"Nota não autorizada (cStat {cstat})", "Nota ignorada")); continue

        emit, dest = txt(inf, "n:emit/n:CNPJ"), txt(inf, "n:dest/n:CNPJ")
        tpnf, fin = txt(inf, "n:ide/n:tpNF"), txt(inf, "n:ide/n:finNFe")
        if emit == cnpj and tpnf == "1": tipo = "Venda"
        elif emit == cnpj and tpnf == "0" and fin == "4": tipo = "Devolução"
        elif dest == cnpj: tipo = "Compra"
        else:
            ocorrencias.append((nome, chave, "Nota não envolve a empresa analisada", "Nota ignorada")); continue

        data = pd.to_datetime(txt(inf, "n:ide/n:dhEmi")[:10])
        uf = txt(inf, "n:dest/n:enderDest/n:UF") if tipo != "Compra" else txt(inf, "n:emit/n:enderEmit/n:UF")
        soma_itens = 0.0
        for det in inf.findall("n:det", NS):
            vprod, vdesc = num(det, "n:prod/n:vProd"), num(det, "n:prod/n:vDesc")
            soma_itens += vprod - vdesc
            itens.append({
                "chave": chave, "data": data, "mes": data.month, "tipo": tipo, "uf": uf,
                "cfop": txt(det, "n:prod/n:CFOP"), "cprod": txt(det, "n:prod/n:cProd"),
                "produto": txt(det, "n:prod/n:xProd"), "ncm": txt(det, "n:prod/n:NCM"),
                "qtd": num(det, "n:prod/n:qCom"), "vprod": vprod, "vdesc": vdesc,
                "icms": num(det, ".//n:ICMS//n:vICMS"), "pis": num(det, ".//n:PIS//n:vPIS"),
                "cofins": num(det, ".//n:COFINS//n:vCOFINS"),
            })
        vnf = num(inf, "n:total/n:ICMSTot/n:vNF")
        if abs(soma_itens - vnf) > 0.05:
            ocorrencias.append((nome, chave, f"Soma dos itens ({soma_itens:.2f}) ≠ total da nota ({vnf:.2f})", "Mantida — revisar"))
    return pd.DataFrame(itens), ocorrencias, len(arquivos)

# ---------------------------------------------------------------- cálculo
def preparar(df):
    compras = df[df.tipo == "Compra"]
    # custo médio ponderado anual: valor pago menos ICMS recuperável (Lucro Presumido não credita PIS/COFINS)
    custo = ((compras.vprod - compras.vdesc - compras.icms).groupby(compras.cprod).sum()
             / compras.qtd.groupby(compras.cprod).sum()).rename("custo_unit")
    df = df.merge(custo, left_on="cprod", how="left", right_index=True)
    sinal = df.tipo.map({"Venda": 1, "Devolução": -1, "Compra": 0})
    df["receita_bruta"] = (df.vprod - df.vdesc) * sinal
    for c in ("icms", "pis", "cofins"): df[c + "_venda"] = df[c] * sinal
    df["cmv"] = df.qtd * df.custo_unit * sinal
    df["valor_compra"] = (df.vprod - df.vdesc) * (df.tipo == "Compra")
    return df

# ---------------------------------------------------------------- relatório
F = "Arial"
NAVY = PatternFill("solid", fgColor="1F3864"); SUB = PatternFill("solid", fgColor="D9E1F2")
GREEN = PatternFill("solid", fgColor="C6EFCE"); RED = PatternFill("solid", fgColor="FFC7CE")
HDR = Font(name=F, bold=True, color="FFFFFF"); BOLD = Font(name=F, bold=True); BLK = Font(name=F)
MON = 'R$ #,##0;(R$ #,##0);-'; PCT = '0.0%;(0.0%);-'; QTD = '#,##0'
thin = Side(style="thin", color="BFBFBF")

def titulo(ws, texto, sub, largura):
    ws["A1"] = texto; ws["A1"].font = Font(name=F, bold=True, size=14, color="FFFFFF")
    ws["A2"] = sub; ws["A2"].font = Font(name=F, italic=True, color="FFFFFF")
    for r in (1, 2):
        for c in range(1, largura + 1): ws.cell(row=r, column=c).fill = NAVY

def cabecalho(ws, r, nomes):
    for i, n in enumerate(nomes, 1):
        c = ws.cell(row=r, column=i, value=n); c.font = HDR; c.fill = NAVY
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30

def gerar_relatorio(df, ocorrencias, n_arquivos, empresa, saida):
    wb = Workbook()
    resumo = wb.active; resumo.title = "Resumo"
    dre = wb.create_sheet("DRE Mensal"); prod = wb.create_sheet("Margem por Produto")
    qual = wb.create_sheet("Qualidade dos Dados"); base = wb.create_sheet("Base Itens")

    # ---- Base Itens (dados brutos — tudo o resto é fórmula em cima dela)
    cols = [("chave", "Chave de acesso"), ("data", "Data"), ("mes", "Mês"), ("tipo", "Tipo"), ("uf", "UF"),
            ("cfop", "CFOP"), ("cprod", "Código"), ("produto", "Produto"), ("ncm", "NCM"), ("qtd", "Qtd"),
            ("receita_bruta", "Receita bruta"), ("icms_venda", "ICMS"), ("pis_venda", "PIS"),
            ("cofins_venda", "COFINS"), ("custo_unit", "Custo médio unit."), ("cmv", "CMV"), ("valor_compra", "Valor compra")]
    cabecalho(base, 1, [c[1] for c in cols])
    df = df.sort_values(["data", "chave"])
    for r, row in enumerate(df[[c[0] for c in cols]].itertuples(index=False), 2):
        for i, v in enumerate(row, 1):
            cell = base.cell(row=r, column=i, value=v.to_pydatetime() if hasattr(v, "to_pydatetime") else v)
            cell.font = BLK
    n = len(df) + 1
    fmts = {2: "dd/mm/yyyy", 10: QTD, 11: MON, 12: MON, 13: MON, 14: MON, 15: 'R$ #,##0.00', 16: MON, 17: MON}
    for col, fmt in fmts.items():
        for r in range(2, n + 1): base.cell(row=r, column=col).number_format = fmt
    for i, w in enumerate([46, 11, 6, 11, 5, 7, 8, 34, 10, 8, 14, 12, 10, 11, 14, 13, 14], 1):
        base.column_dimensions[get_column_letter(i)].width = w
    base.freeze_panes = "A2"; base.auto_filter.ref = f"A1:Q{n}"
    R = lambda col: f"'Base Itens'!${col}$2:${col}${n}"
    MES, TIPO, COD, QT, RB, IC, PI, CO, CM, VC = (R(c) for c in "CDGJKLMNPQ")

    # ---- DRE Mensal
    titulo(dre, "DRE gerencial mensal — 2025", f"{empresa} · valores em R$ · gerada a partir das NF-e", 14)
    cabecalho(dre, 4, ["Linha"] + MESES + ["Total 2025"])
    linhas = [
        (5, "Receita bruta de vendas", lambda m: f'=SUMIFS({RB},{MES},{m},{TIPO},"Venda")'),
        (6, "(−) Devoluções de venda", lambda m: f'=SUMIFS({RB},{MES},{m},{TIPO},"Devolução")'),
        (7, "(−) ICMS", lambda m: f"=-SUMIFS({IC},{MES},{m})"),
        (8, "(−) PIS", lambda m: f"=-SUMIFS({PI},{MES},{m})"),
        (9, "(−) COFINS", lambda m: f"=-SUMIFS({CO},{MES},{m})"),
        (10, "Receita líquida", lambda m, c=None: None),
        (11, "(−) CMV", lambda m: f"=-SUMIFS({CM},{MES},{m})"),
        (12, "Lucro bruto", None), (13, "Margem bruta", None),
        (15, "Compras do mês (informativo)", lambda m: f'=SUMIFS({VC},{MES},{m},{TIPO},"Compra")'),
    ]
    for r, nome, fn in linhas:
        dre.cell(row=r, column=1, value=nome).font = BOLD if r in (10, 12, 13) else BLK
        for m in range(1, 13):
            col = get_column_letter(m + 1); c = dre.cell(row=r, column=m + 1)
            if r == 10: c.value = f"=SUM({col}5:{col}9)"
            elif r == 12: c.value = f"={col}10+{col}11"
            elif r == 13: c.value = f"=IF({col}10=0,0,{col}12/{col}10)"
            else: c.value = fn(m)
        t = dre.cell(row=r, column=14)
        t.value = "=IF(N10=0,0,N12/N10)" if r == 13 else f"=SUM(B{r}:M{r})"
        for c in range(2, 15):
            cell = dre.cell(row=r, column=c); cell.number_format = PCT if r == 13 else MON
            cell.font = BOLD if r in (10, 12, 13) or c == 14 else BLK
            if r in (10, 12): cell.border = Border(top=thin)
    for c in range(1, 15): dre.cell(row=12, column=c).fill = SUB
    dre.column_dimensions["A"].width = 30
    for c in range(2, 15): dre.column_dimensions[get_column_letter(c)].width = 13
    dre.freeze_panes = "B5"

    # ---- Margem por Produto (ordem = curva ABC por receita líquida)
    vendas = df[df.tipo != "Compra"].copy()
    vendas["rl"] = vendas.receita_bruta - vendas.icms_venda - vendas.pis_venda - vendas.cofins_venda
    ordem = vendas.groupby(["cprod", "produto"]).rl.sum().sort_values(ascending=False).reset_index()
    titulo(prod, "Margem e curva ABC por produto — 2025", "Ordenado por receita líquida · classe A até 80%, B até 95%, C o restante", 11)
    cabecalho(prod, 4, ["Código", "Produto", "Qtd vendida", "Receita líquida", "CMV", "Lucro bruto",
                        "Margem bruta", "Preço médio líq.", "Participação", "% acumulado", "Classe ABC"])
    ini = 5
    for i, row in enumerate(ordem.itertuples(index=False)):
        r = ini + i
        prod.cell(row=r, column=1, value=row.cprod); prod.cell(row=r, column=2, value=row.produto)
        prod[f"C{r}"] = f'=SUMIFS({QT},{COD},A{r},{TIPO},"Venda")-SUMIFS({QT},{COD},A{r},{TIPO},"Devolução")'
        prod[f"D{r}"] = f"=SUMIFS({RB},{COD},A{r})-SUMIFS({IC},{COD},A{r})-SUMIFS({PI},{COD},A{r})-SUMIFS({CO},{COD},A{r})"
        prod[f"E{r}"] = f"=SUMIFS({CM},{COD},A{r})"
        prod[f"F{r}"] = f"=D{r}-E{r}"
        prod[f"G{r}"] = f"=IF(D{r}=0,0,F{r}/D{r})"
        prod[f"H{r}"] = f"=IF(C{r}=0,0,D{r}/C{r})"
        fim = ini + len(ordem) - 1
        prod[f"I{r}"] = f"=D{r}/SUM($D${ini}:$D${fim})"
        prod[f"J{r}"] = f"=SUM($I${ini}:I{r})"
        prod[f"K{r}"] = f'=IF(J{r}-I{r}<0.8,"A",IF(J{r}-I{r}<0.95,"B","C"))'
        for c, fmt in zip("CDEFGHIJ", [QTD, MON, MON, MON, PCT, 'R$ #,##0.00', PCT, PCT]):
            prod[f"{c}{r}"].number_format = fmt
        for c in range(1, 12): prod.cell(row=r, column=c).font = BLK
    tot = fim + 1
    prod.cell(row=tot, column=2, value="Total").font = BOLD
    for c in "CDEF": prod[f"{c}{tot}"] = f"=SUM({c}{ini}:{c}{fim})"; prod[f"{c}{tot}"].font = BOLD
    prod[f"G{tot}"] = f"=IF(D{tot}=0,0,F{tot}/D{tot})"; prod[f"G{tot}"].font = BOLD
    for c, fmt in zip("CDEFG", [QTD, MON, MON, MON, PCT]):
        prod[f"{c}{tot}"].number_format = fmt; prod[f"{c}{tot}"].border = Border(top=thin)
    prod.conditional_formatting.add(f"G{ini}:G{fim}", FormulaRule(formula=[f"G{ini}<0.15"], fill=RED))
    prod.conditional_formatting.add(f"G{ini}:G{fim}", FormulaRule(formula=[f"G{ini}>=0.3"], fill=GREEN))
    prod[f"A{tot+2}"] = "Vermelho: margem bruta abaixo de 15%. Verde: 30% ou mais."
    prod[f"A{tot+2}"].font = Font(name=F, italic=True, color="595959")
    for c, w in zip("ABCDEFGHIJK", [9, 36, 12, 15, 15, 15, 12, 14, 12, 12, 10]):
        prod.column_dimensions[c].width = w

    # ---- Qualidade dos Dados
    titulo(qual, "Qualidade dos dados", "Checagens automáticas feitas antes da análise", 4)
    notas = df.drop_duplicates("chave").tipo.value_counts()
    resumo_q = [("Arquivos XML lidos", n_arquivos), ("Notas válidas analisadas", df.chave.nunique()),
                ("  Vendas", int(notas.get("Venda", 0))), ("  Devoluções de venda", int(notas.get("Devolução", 0))),
                ("  Compras", int(notas.get("Compra", 0))), ("Itens de nota processados", len(df)),
                ("Ocorrências encontradas", len(ocorrencias))]
    for i, (k, v) in enumerate(resumo_q, 4):
        qual.cell(row=i, column=1, value=k).font = BLK; qual.cell(row=i, column=2, value=v).font = BOLD
    r0 = 5 + len(resumo_q)
    cabecalho(qual, r0, ["Arquivo", "Chave", "Ocorrência", "Tratamento"])
    for i, oc in enumerate(ocorrencias, r0 + 1):
        for j, v in enumerate(oc, 1): qual.cell(row=i, column=j, value=v).font = BLK
    if not ocorrencias: qual.cell(row=r0 + 1, column=1, value="Nenhuma ocorrência.").font = BLK
    for c, w in zip("ABCD", [58, 48, 52, 40]): qual.column_dimensions[c].width = w

    # ---- Resumo
    resumo.sheet_view.showGridLines = False
    titulo(resumo, "Análise de margem a partir das NF-e — 2025", f"{empresa} (empresa fictícia) · Lucro Presumido", 12)
    kpis = [("B", "Receita líquida", "='DRE Mensal'!N10", MON), ("E", "Lucro bruto", "='DRE Mensal'!N12", MON),
            ("H", "Margem bruta", "='DRE Mensal'!N13", PCT),
            ("K", "Devoluções", "=-'DRE Mensal'!N6/'DRE Mensal'!N5", '0.0%" da receita"')]
    for col, lab, f, fmt in kpis:
        c2 = get_column_letter(resumo[col + "4"].column + 1)
        resumo.merge_cells(f"{col}4:{c2}4"); resumo.merge_cells(f"{col}5:{c2}6")
        resumo[f"{col}4"] = lab; resumo[f"{col}4"].font = Font(name=F, size=9, color="595959")
        resumo[f"{col}5"] = f; resumo[f"{col}5"].number_format = fmt
        resumo[f"{col}5"].font = Font(name=F, bold=True, size=16, color="1F3864")
        for cc in (col, c2):
            for rr in (4, 5, 6): resumo[f"{cc}{rr}"].fill = PatternFill("solid", fgColor="EEF3FA")
        resumo[f"{col}4"].alignment = resumo[f"{col}5"].alignment = Alignment(horizontal="center", vertical="center")
    menor = f"'Margem por Produto'!$G${ini}:$G${fim}"
    resumo["B8"] = (f'="Produto com menor margem: "&INDEX(\'Margem por Produto\'!$B${ini}:$B${fim},MATCH(MIN({menor}),{menor},0))'
                    f'&" ("&TEXT(MIN({menor}),"0%")&"). Produtos classe A: "&COUNTIF(\'Margem por Produto\'!$K${ini}:$K${fim},"A")'
                    f'&" de "&COUNTA(\'Margem por Produto\'!$A${ini}:$A${fim})&"."')
    resumo["B8"].font = Font(name=F, color="1F3864")
    for c in "ABCDEFGHIJKLM": resumo.column_dimensions[c].width = 11
    resumo.column_dimensions["A"].width = 2

    g1 = LineChart(); g1.title = "Receita líquida e lucro bruto por mês (R$)"; g1.height = 8; g1.width = 17
    for r in (10, 12): g1.add_data(Reference(dre, min_col=1, max_col=13, min_row=r), from_rows=True, titles_from_data=True)
    g1.set_categories(Reference(dre, min_col=2, max_col=13, min_row=4))
    g1.y_axis.number_format = "#,##0"; g1.x_axis.delete = g1.y_axis.delete = False
    resumo.add_chart(g1, "B10")
    g2 = BarChart(); g2.type = "bar"; g2.title = "Margem bruta por produto"; g2.height = 8; g2.width = 15
    g2.add_data(Reference(prod, min_col=7, min_row=4, max_row=fim), titles_from_data=True)
    g2.set_categories(Reference(prod, min_col=2, min_row=ini, max_row=fim)); g2.legend = None
    g2.x_axis.delete = g2.y_axis.delete = False; g2.y_axis.number_format = "0%"
    resumo.add_chart(g2, "H10")
    g3 = BarChart(); g3.title = "Lucro bruto por produto (R$)"; g3.height = 8; g3.width = 32
    g3.add_data(Reference(prod, min_col=6, min_row=4, max_row=fim), titles_from_data=True)
    g3.set_categories(Reference(prod, min_col=1, min_row=ini, max_row=fim)); g3.legend = None
    g3.y_axis.number_format = "#,##0"; g3.x_axis.delete = g3.y_axis.delete = False
    resumo.add_chart(g3, "B27")
    wb.save(saida)

def main():
    ap = argparse.ArgumentParser(description="DRE e margem por produto a partir de XMLs de NF-e")
    ap.add_argument("--pasta", default="xmls"); ap.add_argument("--cnpj", default="11222333000181")
    ap.add_argument("--empresa", default="Atlas Distribuidora de Suprimentos Industriais Ltda")
    ap.add_argument("--saida", default="relatorio_margem_2025.xlsx")
    a = ap.parse_args()
    df, oc, n = ler_xmls(a.pasta, a.cnpj)
    if df.empty: raise SystemExit("Nenhuma nota válida encontrada.")
    df = preparar(df)
    gerar_relatorio(df, oc, n, a.empresa, a.saida)
    print(f"{n} arquivos lidos · {df.chave.nunique()} notas válidas · {len(oc)} ocorrências · relatório: {a.saida}")

if __name__ == "__main__":
    main()
